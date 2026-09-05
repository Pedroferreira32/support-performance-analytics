from __future__ import annotations

import json
import io
import tempfile
import threading
import unittest
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import motor_premiacao as core
import servidor_html as web


class ServidorHtmlTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        web.DEMO_MODE = False
        web.DB_PATH = Path(self.temp.name) / "historico_teste.db"
        db = web.banco()
        dados = pd.DataFrame(
            [
                {
                    "Protocolo": "1",
                    "User ID": "Lucas Rocha",
                    "Filas": "Suporte",
                    "Filas Transfers": "",
                    "Iniciado": "03/08/2026 10:00",
                    "Fim": "03/08/2026 11:00",
                    "Rating": 5,
                },
                {
                    "Protocolo": "2",
                    "User ID": "Agente Demonstração Excluído",
                    "Filas": "Suporte",
                    "Filas Transfers": "",
                    "Iniciado": "03/08/2026 10:00",
                    "Fim": "03/08/2026 10:06",
                    "Rating": 5,
                },
                {
                    "Protocolo": "3",
                    "User ID": "Lucas Rocha",
                    "Filas": "Suporte",
                    "Filas Transfers": "",
                    "Iniciado": "03/08/2026 10:00",
                    "Fim": "04/08/2026 06:00",
                    "Rating": 4,
                },
            ]
        )
        resultado = core.processar_dados(
            dados, core.configuracao_oficial("2026-08"), db, origem="teste.xlsx"
        )
        db.salvar(resultado)

        self.server = web.ThreadingHTTPServer(("127.0.0.1", 0), web.PainelHandler)
        self.port = self.server.server_address[1]
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    def tearDown(self) -> None:
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=3)
        self.temp.cleanup()

    def url(self, caminho: str) -> str:
        return f"http://127.0.0.1:{self.port}{caminho}"

    def get_json(self, caminho: str) -> dict:
        with urllib.request.urlopen(self.url(caminho), timeout=5) as resposta:
            self.assertEqual(resposta.status, 200)
            return json.loads(resposta.read().decode("utf-8"))

    def test_dashboard_e_apis_principais(self) -> None:
        with urllib.request.urlopen(self.url("/"), timeout=5) as resposta:
            html = resposta.read().decode("utf-8")
        self.assertIn("Premiação do Suporte", html)
        self.assertIn("meta mínima: 85 pontos", html)
        self.assertIn("Premiados", html)
        self.assertIn('id="powerpoint-download" disabled', html)
        self.assertIn('id="view-automaticos"', html)
        self.assertIn('id="view-gerencial"', html)
        self.assertIn("Análise gerencial", html)
        self.assertEqual(html.count('<script src="'), 1)
        self.assertNotIn('src="/powerpoint.js"', html)
        self.assertNotIn('src="/vendor/pptxgen.min.js"', html)

        with urllib.request.urlopen(self.url("/dashboard.js"), timeout=5) as resposta:
            javascript = resposta.read().decode("utf-8")
        self.assertIn("Dashboard local autocontido", javascript)
        self.assertIn("var PptxGenJS=", javascript)
        self.assertIn("root.PremiacaoPowerPoint =", javascript)
        self.assertIn("function montarResumoPowerPoint", javascript)
        self.assertIn("META_ELEGIBILIDADE = 85", javascript)
        self.assertIn("31,50 pontos fixos", javascript)
        self.assertIn("function renderGerencial", javascript)
        self.assertIn("/api/gerencial", javascript)
        self.assertNotIn("/api/powerpoint-resumo?", javascript)

        resumo = self.get_json("/api/resumo?competencia=2026-08")
        self.assertEqual(resumo["competencia"], "2026-08")
        self.assertEqual(len(resumo["ranking"]), 1)
        self.assertEqual(resumo["ranking"][0]["atendente"], "Lucas Rocha")
        self.assertEqual(resumo["ranking"][0]["premiado"], 1)
        self.assertEqual(resumo["info"]["configuracao"]["nota_minima"], 85.0)
        self.assertEqual(resumo["info"]["estatisticas"]["ATENDENTE_EXCLUIDO"], 1)
        self.assertEqual(resumo["info"]["estatisticas"]["AUTOMATICOS_RECUPERADOS"], 1)

        gerencial = self.get_json("/api/gerencial?competencia=2026-08")
        self.assertEqual(gerencial["competencia"], "2026-08")
        self.assertEqual(gerencial["metricas"]["funcionarios"], 1)
        self.assertEqual(gerencial["metricas"]["premiados"], 1)
        self.assertEqual(gerencial["metricas"]["elegiveis"], 1)
        self.assertAlmostEqual(gerencial["metricas"]["taxa_validacao"], 200 / 3)
        self.assertFalse(gerencial["comparacao_disponivel"])
        self.assertFalse(gerencial["comparacao_confiavel"])
        componentes = {item["indicador"]: item for item in gerencial["componentes"]}
        self.assertTrue(componentes["Tempo Total"]["protegido"])
        self.assertTrue(componentes["TMA"]["protegido"])
        self.assertEqual(gerencial["acoes"][0]["situacao"], "Premiado")

        automaticos = self.get_json("/api/automaticos?competencia=2026-08")
        self.assertFalse(automaticos["requer_reprocessamento"])
        self.assertEqual(automaticos["metricas"]["identificados"], 1)
        self.assertEqual(automaticos["metricas"]["incluidos"], 1)
        self.assertEqual(automaticos["metricas"]["recuperados"], 1)
        self.assertEqual(automaticos["registros"][0]["protocolo"], "3")
        self.assertIn("Tempo/TMA fixos", automaticos["registros"][0]["tratamento"])

        auditoria = self.get_json("/api/auditoria?competencia=2026-08")
        self.assertEqual(len(auditoria["exclusoes"]), 1)
        self.assertEqual(auditoria["exclusoes"][0]["motivo_codigo"], "ATENDENTE_EXCLUIDO")

        powerpoint = self.get_json("/api/powerpoint-resumo?competencia=2026-08")
        self.assertEqual(powerpoint["competencia"], "2026-08")
        self.assertEqual(powerpoint["quantidade_slides"], 13)
        self.assertEqual(powerpoint["atual"]["ranking"][0]["atendente"], "Lucas Rocha")
        self.assertEqual(sum(item["slides"] for item in powerpoint["roteiro"]), 13)

        payload = json.dumps(
            {
                "competencia": "2026-08",
                "atendente": "Lucas Rocha",
                "feedback": "Feedback revisado no dashboard HTML.",
            }
        ).encode("utf-8")
        request = urllib.request.Request(
            self.url("/api/feedback"),
            data=payload,
            method="POST",
            headers={"Content-Type": "application/json"},
        )
        with urllib.request.urlopen(request, timeout=5) as resposta:
            self.assertEqual(resposta.status, 200)
        historico = self.get_json(
            "/api/historico?atendente=" + urllib.parse.quote("Lucas Rocha")
        )
        self.assertEqual(historico["historico"][0]["feedback"], "Feedback revisado no dashboard HTML.")

        with urllib.request.urlopen(
            self.url("/api/exportar?competencia=2026-08"), timeout=10
        ) as resposta:
            self.assertEqual(resposta.status, 200)
            conteudo_excel = resposta.read()
            self.assertGreater(len(conteudo_excel), 5000)
        workbook = load_workbook(io.BytesIO(conteudo_excel), read_only=True)
        self.assertIn("Finalizados_Automaticos", workbook.sheetnames)

    def test_automaticos_avisa_quando_agosto_precisa_ser_reprocessado(self) -> None:
        db = web.banco()
        info = db.competencia_info("2026-08")
        configuracao = json.loads(info["configuracao_json"])
        configuracao.pop("pontuacao_tempo_tma_fixa", None)
        with db.conectar() as con:
            con.execute(
                "UPDATE competencias SET configuracao_json = ? WHERE competencia = ?",
                (json.dumps(configuracao, ensure_ascii=False), "2026-08"),
            )

        automaticos = self.get_json("/api/automaticos?competencia=2026-08")
        self.assertTrue(automaticos["requer_reprocessamento"])

    def test_gerencial_compara_competencias_do_mesmo_perfil(self) -> None:
        db = web.banco()
        setembro = pd.DataFrame(
            [
                {
                    "Protocolo": "10",
                    "User ID": "Lucas Rocha",
                    "Filas": "Suporte",
                    "Filas Transfers": "",
                    "Iniciado": "02/09/2026 10:00",
                    "Fim": "02/09/2026 11:00",
                    "Rating": 5,
                }
            ]
        )
        db.salvar(
            core.processar_dados(
                setembro,
                core.configuracao_oficial("2026-09"),
                db,
                origem="setembro.xlsx",
            )
        )

        gerencial = self.get_json("/api/gerencial?competencia=2026-09")
        self.assertTrue(gerencial["comparacao_disponivel"])
        self.assertTrue(gerencial["comparacao_confiavel"])
        self.assertEqual(gerencial["competencia_anterior"], "2026-08")
        self.assertEqual(len(gerencial["serie_equipe"]), 2)
        self.assertEqual(gerencial["comparativo"][0]["atendente"], "Lucas Rocha")
        self.assertIsNotNone(gerencial["comparativo"][0]["delta_nota"])

    def test_upload_julho_aplica_perfil_automatico(self) -> None:
        csv = (
            "Protocolo;User ID;Filas;Filas Transfers;Iniciado;Fim;Rating\n"
            "10;Lucas Rocha;Suporte;;05/07/2026 05:30;05/07/2026 06:00;5\n"
        ).encode("utf-8")
        query = urllib.parse.urlencode(
            {
                "competencia": "07/2026",
                "arquivo": "julho.csv",
                "excluidos": core.ATENDENTES_FORA_CAMPANHA,
            }
        )
        request = urllib.request.Request(
            self.url(f"/api/processar?{query}"),
            data=csv,
            method="POST",
            headers={"Content-Type": "application/octet-stream"},
        )
        with urllib.request.urlopen(request, timeout=10) as resposta:
            payload = json.loads(resposta.read().decode("utf-8"))
        self.assertTrue(payload["ok"])
        self.assertEqual(payload["validos"], 1)
        self.assertEqual(payload["estatisticas"]["AUTOMATICOS_VALIDOS"], 1)

        resumo = self.get_json("/api/resumo?competencia=2026-07")
        self.assertEqual(resumo["info"]["modo"], "neutralizado")
        self.assertAlmostEqual(resumo["ranking"][0]["pontos_tempo"], 7.88)
        self.assertAlmostEqual(resumo["ranking"][0]["pontos_tma"], 23.62)

        powerpoint = self.get_json("/api/powerpoint-resumo?competencia=2026-07")
        self.assertEqual(powerpoint["atual"]["info"]["modo"], "neutralizado")
        self.assertEqual(powerpoint["anterior"], None)

    def test_upload_rejeita_competencia_divergente(self) -> None:
        csv = (
            "Protocolo;User ID;Setores;Setores Transfers;Iniciado;Fim;Rating\n"
            "20;Lucas Rocha;Suporte;;03/08/2026 10:00;03/08/2026 11:00;5\n"
        ).encode("utf-8")
        query = urllib.parse.urlencode(
            {
                "competencia": "09/2026",
                "arquivo": "agosto.csv",
                "excluidos": core.ATENDENTES_FORA_CAMPANHA,
            }
        )
        request = urllib.request.Request(
            self.url(f"/api/processar?{query}"),
            data=csv,
            method="POST",
            headers={"Content-Type": "application/octet-stream"},
        )
        with self.assertRaises(urllib.error.HTTPError) as contexto:
            urllib.request.urlopen(request, timeout=10)

        self.assertEqual(contexto.exception.code, 400)
        payload = json.loads(contexto.exception.read().decode("utf-8"))
        self.assertIn("08/2026", payload["erro"])
        self.assertIn("09/2026", payload["erro"])


if __name__ == "__main__":
    unittest.main()
